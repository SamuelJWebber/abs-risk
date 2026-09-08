"""Build data/ccmr_level.csv from the CFPB Consumer Credit Card Market Report (CCMR) figure-data workbooks.

Run from the repo root: `python -m uv run python -m absrisk.shape.build_level`. Needs the two workbooks under
data/raw/scout/cfpb/ (design/scout-ccmr-tiers.md section 1 has the download URLs and hashes). Every number
written is read from a workbook cell found by its row label and column header, never by position.

What is built, per year and per bureau panel ("series"):

- `gp_chargeoff_rate`: annual mean of the CFPB's annualized gross charge-off rate on general-purpose card
  balances (monthly in the CCIP series, quarterly in the CCP series). `gp_chargeoff_rate_ye` is the
  December / Q4 value. The two panels are not on the same basis (2025 report fn 13; scout note section 3.2)
  and are kept as separate rows, never spliced.
- `share_<tier>`: the share of balances held by each of the six credit tiers. The CCMR does not publish
  balances by tier, so the share is constructed from what it does publish:
    CCIP (2025 workbook, 2014-2024): consumers holding a general-purpose card by tier at year-end 2023
      (Figure 3, held fixed across years) x average per-cardholder cycle-ending balance by tier
      (Figure 16, annual mean of monthly values; all cards, not general purpose only).
    CCP (2023 workbook, 2013-2022): general-purpose accounts by tier at year-end 2022 (Section 5 Figure 17)
      x average per-account general-purpose balance by tier at year-end 2022 (Section 3 Figure 7), which
      is the exact general-purpose balance by tier at that anchor date, moved to other years by the ratio
      of the tier's average per-cardholder balance (Section 3 Figure 5, annual mean of quarterly values)
      to its 2022Q4 value.
  Both are then normalised to sum to one. The tier populations are by *current* score, so a tier's share
  tracks a changing set of people (2025 report PDF p15). design/shape-sources.md discusses the limits.
"""

from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl
import pandas as pd

from .tiers import TIERS

RAW_DIR = Path("data/raw/scout/cfpb")
OUT_PATH = Path("data/ccmr_level.csv")
WB_2025 = "cfpb_consumer-credit-card-market-report-figure-data_2025.xlsm"
WB_2023 = "cfpb_consumer-credit-card-market-report-figure-data_2023.xlsx"

# Column / row labels as they appear in the workbooks -> tier ids.
TIER_LABELS = {
    "Deep subprime": "deep_subprime",
    "Subprime": "subprime",
    "Near-prime": "near_prime",
    "Prime": "prime",
    "Prime plus": "prime_plus",
    "Superprime": "superprime",
}

COLUMNS = [
    "year",
    "series",
    "gp_chargeoff_rate",
    "gp_chargeoff_rate_ye",
    "n_periods",
    *[f"share_{t}" for t in TIERS],
    "mix_periods",
    "citation_rate",
    "citation_mix",
]


def figure_block(ws, label: str) -> pd.DataFrame:
    """The data table under the row whose first cell is `label` (e.g. "Figure 53").

    Every CCMR workbook uses the same layout: label row, title row, header row, data rows, blank row.
    Returns a DataFrame with the header labels (stripped) as columns; the figure title is in `.attrs`.
    """
    rows = list(ws.iter_rows(values_only=True))
    starts = [i for i, r in enumerate(rows) if r and isinstance(r[0], str) and r[0].strip() == label]
    if len(starts) != 1:
        raise ValueError(f"{label!r}: found {len(starts)} label rows in sheet {ws.title!r}")
    i = starts[0]
    title = str(rows[i + 1][0]).strip()
    header = ["" if h is None else str(h).strip() for h in rows[i + 2]]
    ncol = max(j for j, h in enumerate(header) if h) + 1
    data = []
    for r in rows[i + 3 :]:
        if r is None or all(v is None for v in r):
            break
        data.append(list(r[:ncol]))
    df = pd.DataFrame(data, columns=header[:ncol])
    df.attrs["title"] = title
    df.attrs["sheet"] = ws.title
    return df


def _tier_columns(df: pd.DataFrame) -> dict[str, str]:
    """Map tier id -> the DataFrame column holding it; fails if any tier is missing."""
    cols = {c.strip(): c for c in df.columns}
    out = {}
    for label, tier in TIER_LABELS.items():
        if label not in cols:
            raise ValueError(f"tier column {label!r} missing from {df.attrs.get('title')!r}")
        out[tier] = cols[label]
    return out


def _tier_rows(df: pd.DataFrame, key_col: str, value_col: str) -> dict[str, float]:
    keys = {str(k).strip(): v for k, v in zip(df[key_col], df[value_col])}
    out = {}
    for label, tier in TIER_LABELS.items():
        if label not in keys:
            raise ValueError(f"tier row {label!r} missing from {df.attrs.get('title')!r}")
        out[tier] = float(keys[label])
    return out


def _shares(weights: dict[str, float]) -> dict[str, float]:
    """Normalise to six-decimal shares that sum to exactly 1 (rounding residual goes to the largest tier)."""
    total = sum(weights.values())
    shares = {t: round(weights[t] / total, 6) for t in TIERS}
    largest = max(TIERS, key=lambda t: shares[t])
    shares[largest] = round(shares[largest] + (1.0 - sum(shares.values())), 6)
    return shares


def build_ccip(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    co = figure_block(wb["Section 4 - Pmts, Debt, Coll."], "Figure 53")
    if "General purpose" not in co.columns or "Date" not in co.columns:
        raise ValueError(f"unexpected columns in Figure 53: {list(co.columns)}")
    co["date"] = pd.to_datetime(co["Date"])
    co = co.sort_values("date")
    co["year"] = co["date"].dt.year

    consumers = figure_block(wb["Section 2 - Use of credit"], "Figure 3")
    n_by_tier = _tier_rows(consumers, "Credit score tier", "General purpose")

    bal = figure_block(wb["Section 2 - Use of credit"], "Figure 16")
    bal["date"] = pd.to_datetime(bal["Date"])
    bal = bal.sort_values("date")
    bal["year"] = bal["date"].dt.year
    bcols = _tier_columns(bal)

    cite_rate = (
        f"CCMR 2025 workbook {WB_2025}, sheet 'Section 4 - Pmts, Debt, Coll.', Figure 53 "
        f"'{co.attrs['title']}', column 'General purpose'; annual mean of the monthly values, "
        "year-end = December"
    )
    cite_mix = (
        f"CCMR 2025 workbook {WB_2025}, sheet 'Section 2 - Use of credit': Figure 3 "
        f"'{consumers.attrs['title']}' column 'General purpose' (consumers by tier, held fixed for all "
        f"years) x Figure 16 '{bal.attrs['title']}' (annual mean of monthly per-cardholder balance by "
        "tier); share = product / sum over tiers"
    )

    out = []
    for year, g in co.groupby("year"):
        b = bal[bal["year"] == year]
        if b.empty:
            raise ValueError(f"no per-cardholder balances for {year}")
        weights = {t: n_by_tier[t] * float(b[bcols[t]].mean()) for t in TIERS}
        months = b["date"].dt.strftime("%b").tolist()
        row = {
            "year": int(year),
            "series": "ccip",
            "gp_chargeoff_rate": float(g["General purpose"].mean()),
            "gp_chargeoff_rate_ye": float(g.iloc[-1]["General purpose"]),
            "n_periods": int(len(g)),
            "mix_periods": f"{len(b)} months {months[0]}-{months[-1]} {year}",
            "citation_rate": cite_rate,
            "citation_mix": cite_mix,
        }
        row.update({f"share_{t}": v for t, v in _shares(weights).items()})
        out.append(row)
    return out


def build_ccp(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    use = wb["Section 3 - Use of credit"]
    avail = wb["Section 5 - Avail of Credit"]

    co = figure_block(use, "Figure 18")
    if "General purpose" not in co.columns or "Quarter" not in co.columns:
        raise ValueError(f"unexpected columns in Section 3 Figure 18: {list(co.columns)}")
    co["year"] = co["Quarter"].astype(str).str[:4].astype(int)
    co = co.sort_values("Quarter")

    accts = figure_block(avail, "Figure 17")
    acct_gp = accts[accts.iloc[:, 0].astype(str).str.strip() == "General purpose"]
    if len(acct_gp) != 1:
        raise ValueError("Section 5 Figure 17: no single 'General purpose' row")
    acols = _tier_columns(accts)
    n_by_tier = {t: float(acct_gp.iloc[0][acols[t]]) for t in TIERS}

    peracct = figure_block(use, "Figure 7")
    pa_gp = peracct[peracct.iloc[:, 0].astype(str).str.strip() == "General purpose"]
    if len(pa_gp) != 1:
        raise ValueError("Section 3 Figure 7: no single 'General purpose' row")
    pcols = _tier_columns(peracct)
    bal_anchor = {t: n_by_tier[t] * float(pa_gp.iloc[0][pcols[t]]) for t in TIERS}

    pch = figure_block(use, "Figure 5")
    pch["year"] = pch["Quarter"].astype(str).str[:4].astype(int)
    ccols = _tier_columns(pch)
    anchor_q = pch[pch["Quarter"].astype(str).str.strip() == "2022Q4"]
    if len(anchor_q) != 1:
        raise ValueError("Section 3 Figure 5: no 2022Q4 row")
    pch_anchor = {t: float(anchor_q.iloc[0][ccols[t]]) for t in TIERS}

    cite_rate = (
        f"CCMR 2023 workbook {WB_2023}, sheet 'Section 3 - Use of credit', Figure 18 "
        f"'{co.attrs['title']}', column 'General purpose'; annual mean of the quarterly values, "
        "year-end = Q4"
    )
    cite_mix = (
        f"CCMR 2023 workbook {WB_2023}: sheet 'Section 5 - Avail of Credit' Figure 17 "
        f"'{accts.attrs['title']}' row 'General purpose' x sheet 'Section 3 - Use of credit' Figure 7 "
        f"'{peracct.attrs['title']}' row 'General purpose' (= general-purpose balances by tier at "
        f"year-end 2022), scaled by Figure 5 '{pch.attrs['title']}' annual mean / 2022Q4 value by tier; "
        "share = product / sum over tiers"
    )

    out = []
    for year, g in co.groupby("year"):
        b = pch[pch["year"] == year]
        if b.empty:
            raise ValueError(f"no per-cardholder balances for {year}")
        weights = {t: bal_anchor[t] * float(b[ccols[t]].mean()) / pch_anchor[t] for t in TIERS}
        quarters = b["Quarter"].astype(str).tolist()
        row = {
            "year": int(year),
            "series": "ccp",
            "gp_chargeoff_rate": float(g["General purpose"].mean()),
            "gp_chargeoff_rate_ye": float(g.iloc[-1]["General purpose"]),
            "n_periods": int(len(g)),
            "mix_periods": f"{len(b)} quarters {quarters[0]}-{quarters[-1]}",
            "citation_rate": cite_rate,
            "citation_mix": cite_mix,
        }
        row.update({f"share_{t}": v for t, v in _shares(weights).items()})
        out.append(row)
    return out


def build(raw_dir: Path = RAW_DIR) -> pd.DataFrame:
    rows = build_ccip(raw_dir / WB_2025) + build_ccp(raw_dir / WB_2023)
    df = pd.DataFrame(rows, columns=COLUMNS).sort_values(["series", "year"]).reset_index(drop=True)
    num = ["gp_chargeoff_rate", "gp_chargeoff_rate_ye", *[f"share_{t}" for t in TIERS]]
    df[num] = df[num].round(6)
    return df


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    p.add_argument("--raw-dir", type=Path, default=RAW_DIR)
    p.add_argument("--out", type=Path, default=OUT_PATH)
    a = p.parse_args(argv)
    df = build(a.raw_dir)
    a.out.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(a.out, index=False, lineterminator="\n")
    print(f"wrote {a.out} ({len(df)} rows)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
